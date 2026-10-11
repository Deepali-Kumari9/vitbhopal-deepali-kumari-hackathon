"""Build docs/presentation.pptx (7 slides). Export to PDF with LibreOffice:
    soffice --headless --convert-to pdf docs/presentation.pptx --outdir docs
"""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
NAVY, TEAL, RED, WHITE = RGBColor(0x0B, 0x1F, 0x3A), RGBColor(0x2A, 0x9D, 0x8F), RGBColor(0xE6, 0x39, 0x46), RGBColor(255, 255, 255)
INK, MUTED, LIGHT = RGBColor(0x1B, 0x26, 0x3B), RGBColor(0x64, 0x74, 0x8B), RGBColor(0xF1, 0xF5, 0xF9)
ev = json.loads((DOCS / "evaluation_results.json").read_text())

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    return s


def text(slide, x, y, w, h, content, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(8)
        r = p.add_run(); r.text = line
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, "Calibri"
    return tb


def header(slide, title, n):
    rect(slide, 0, 0, 13.333, 1.1, NAVY)
    rect(slide, 0, 1.1, 13.333, 0.06, TEAL)
    text(slide, 0.6, 0.22, 11, 0.7, title, 30, WHITE, True)
    text(slide, 12.2, 7.0, 1, 0.4, f"{n}/7", 12, MUTED, align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, items, size=18):
    return text(slide, x, y, w, 5, ["•  " + i for i in items], size)


def card(slide, x, y, w, h, big, small, color=TEAL):
    rect(slide, x, y, w, h, LIGHT); rect(slide, x, y, 0.08, h, color)
    text(slide, x + 0.25, y + 0.1, w - 0.3, 0.8, big, 30, color, True)
    text(slide, x + 0.25, y + 0.9, w - 0.3, h - 0.9, small, 13, MUTED)


# 1 title
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.333, 7.5, NAVY); rect(s, 0.8, 3.05, 1.2, 0.08, TEAL)
text(s, 0.8, 1.6, 11, 1.2, "RiskPulse AI", 54, WHITE, True)
text(s, 0.8, 3.3, 11, 1, "From unstructured text to risk signals, index weights and stress tests", 24, RGBColor(0xCB, 0xD5, 0xE1))
text(s, 0.8, 5.2, 11, 1.2, ["Deepali Kumari  |  VIT Bhopal University", "S&P Global & Crisil Campus Hackathon 2026"], 18, RGBColor(0x94, 0xA3, 0xB8))

# 2 problem
s = prs.slides.add_slide(blank); header(s, "Problem & Approach", 2)
text(s, 0.6, 1.5, 6, 0.5, "The problem", 22, TEAL, True)
bullets(s, 0.6, 2.2, 5.9, ["Risk news arrives as free text, faster than anyone can read it",
                           "Desks need structured, comparable signals, not headlines",
                           "Those signals must feed real decisions, not just a score"], 22)
text(s, 7.0, 1.5, 6, 0.5, "What I built", 22, TEAL, True)
bullets(s, 7.0, 2.1, 5.8, ["NLP engine: sentiment (-1 to 1), event type, impact (1-10)",
                           "Two sources: news headlines and social posts",
                           "Module A: sentiment-driven index rebalancer",
                           "Module B: event-driven portfolio stress test"], 22)

# 3 architecture
s = prs.slides.add_slide(blank); header(s, "System Design", 3)
s.shapes.add_picture(str(DOCS / "architecture.png"), Inches(0.9), Inches(1.45), width=Inches(11.5))
text(s, 0.6, 6.75, 12, 0.5, "Both modules read the same signal table; the dashboard, API and CSV/JSON export serve it to users.", 14, MUTED)

# 4 implementation
s = prs.slides.add_slide(blank); header(s, "Implementation Highlights", 4)
items = [("Risk engine", "Lexicon sentiment with negation and intensifiers, word-boundary event rules, explainable impact score. Every output lists the words behind it."),
         ("Index rebalancer", "Per-stock sentiment state with half-life decay, exponential tilt, 2%-15% weight caps, turnover tracking."),
         ("Stress tester", "Bonds and loans shocked via duration and spread duration; severity scales the shock; positive news does not trigger it."),
         ("Delivery", "Streamlit dashboard, FastAPI endpoints, CSV/JSON export, 24 automated tests, labelled evaluation set.")]
for i, (t, d) in enumerate(items):
    x, y = 0.6 + (i % 2) * 6.2, 1.6 + (i // 2) * 2.7
    rect(s, x, y, 5.9, 2.4, LIGHT); rect(s, x, y, 5.9, 0.07, TEAL if i < 2 else RED)
    text(s, x + 0.25, y + 0.2, 5.4, 0.5, t, 20, NAVY, True)
    text(s, x + 0.25, y + 0.8, 5.4, 1.5, d, 17, INK)

# 5 results
s = prs.slides.add_slide(blank); header(s, "Key Results", 5)
card(s, 0.6, 1.5, 2.9, 1.6, f"{ev['event_type']['accuracy']:.0%}", f"event-type accuracy (baseline {ev['event_type']['majority_baseline_accuracy']:.0%})")
card(s, 3.7, 1.5, 2.9, 1.6, f"{ev['sentiment']['accuracy']:.0%}", f"sentiment accuracy (baseline {ev['sentiment']['majority_baseline_accuracy']:.0%})")
card(s, 6.8, 1.5, 2.9, 1.6, "~10%", "loss on the $105M book, credit event at impact 8", RED)
card(s, 9.9, 1.5, 2.9, 1.6, "15", "stocks rebalanced across 5 days of signals", RED)
s.shapes.add_picture(str(DOCS / "img" / "weights.png"), Inches(0.5), Inches(3.4), width=Inches(6.1))
s.shapes.add_picture(str(DOCS / "img" / "stress.png"), Inches(6.9), Inches(3.4), width=Inches(6.1))

# 6 impact
s = prs.slides.add_slide(blank); header(s, "Domain Impact", 6)
bullets(s, 0.6, 1.6, 12, ["Turns text into a traceable signal: every score shows the words that produced it",
                          "One signal, two time horizons: fast tactical weights and slower strategic stress losses",
                          "Shows how an event type maps to different damage by asset class (e.g. gov bonds can gain in a geopolitical shock)",
                          "Transparent rules give risk teams something they can audit and challenge before trusting a model",
                          "Modular design: swap the engine for a transformer without touching the downstream modules"], 24)

# 7 limits
s = prs.slides.add_slide(blank); header(s, "Limitations & Next Steps", 7)
text(s, 0.6, 1.5, 6, 0.5, "Limitations", 22, RED, True)
bullets(s, 0.6, 2.1, 5.9, ["Sentiment is a lexicon baseline; misses phrases like \"tensions ease\"",
                           "Impact is a heuristic, not calibrated to returns",
                           "Synthetic data and a small hand-labelled test set",
                           "Shocks ignore correlations, liquidity and Greeks"], 22)
text(s, 7.0, 1.5, 6, 0.5, "Next steps", 22, TEAL, True)
bullets(s, 7.0, 2.1, 5.8, ["Fine-tune a finance transformer (FinBERT-style)",
                           "Larger labelled set and historical backtests",
                           "Calibrate impact against realised moves",
                           "Streaming ingestion and richer stress models"], 22)

prs.save(DOCS / "presentation.pptx")
print("saved", DOCS / "presentation.pptx")
